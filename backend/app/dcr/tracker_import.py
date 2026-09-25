"""Import the existing DCR tracker workbook (SOP 5-A1, sheet "Main Sheet")."""

from pathlib import Path

import openpyxl

from app.dcr.normalize import (
    clean_text,
    norm_category,
    norm_office,
    norm_responsible,
    norm_type,
    norm_yn,
    to_iso_date,
)
from app.dcr.schema import DCR

SHEET = "Main Sheet"
HEADER_CELL = "DCR tracking number"

# Column order of the tracker (A..U)
COLUMNS = [
    "tracking_number",
    "type",
    "critical",
    "occurred_on",
    "entered_by",
    "amex_office",
    "pharma",
    "supplier",
    "customer",
    "project",
    "responsible_party",
    "responsible_party_name",
    "category",
    "description",
    "root_cause",
    "financial_impact",
    "actions_taken",
    "capa_needed",
    "capa_plan",
    "actions",
    "closure_date",
]


def _find_header_row(ws) -> int:
    for row in ws.iter_rows(min_row=1, max_row=30, max_col=1):
        if str(row[0].value or "").strip() == HEADER_CELL:
            return row[0].row
    raise ValueError(f"Header '{HEADER_CELL}' not found in sheet '{SHEET}'")


def import_tracker(path: Path) -> list[DCR]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET]
    header_row = _find_header_row(ws)
    records: list[DCR] = []
    for row_number, values in enumerate(
        ws.iter_rows(min_row=header_row + 1, max_col=len(COLUMNS), values_only=True), start=header_row + 1
    ):
        raw = dict(zip(COLUMNS, values))
        tracking = clean_text(raw["tracking_number"])
        # Pre-numbered but unused rows have only a tracking number.
        if not tracking or all(v is None for k, v in raw.items() if k != "tracking_number"):
            continue

        closure_raw = raw["closure_date"]
        closure = to_iso_date(closure_raw) or clean_text(closure_raw)
        occurred = to_iso_date(raw["occurred_on"]) or clean_text(raw["occurred_on"])

        record = DCR(
            id=f"T-{tracking}",
            tracking_number=tracking,
            type=norm_type(raw["type"]),
            critical=norm_yn(raw["critical"]),
            occurred_on=occurred,
            entered_by=clean_text(raw["entered_by"]),
            amex_office=norm_office(raw["amex_office"]),
            pharma=norm_yn(raw["pharma"]),
            supplier=clean_text(raw["supplier"]),
            customer=clean_text(raw["customer"]),
            project=clean_text(raw["project"]),
            responsible_party=norm_responsible(raw["responsible_party"]),
            responsible_party_name=clean_text(raw["responsible_party_name"]),
            category=norm_category(raw["category"]),
            description=clean_text(raw["description"]),
            root_cause=clean_text(raw["root_cause"]),
            financial_impact=clean_text(raw["financial_impact"]),
            actions_taken=clean_text(raw["actions_taken"]),
            capa_needed=norm_yn(raw["capa_needed"]),
            capa_plan=clean_text(raw["capa_plan"]),
            actions=clean_text(raw["actions"]),
            closure_date=closure,
            source="tracker",
            excel_row=row_number,
        )
        record.status = "Closed" if closure and (to_iso_date(closure) or closure.lower() == "closed") else "Open"
        records.append(record)
    return records
